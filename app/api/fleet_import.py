"""Fleet import — nhập Excel phương tiện/lái xe, theo dõi tiến trình."""
import io
import json
import uuid
from pathlib import Path
from typing import Optional

from fastapi import APIRouter, Depends, UploadFile, File, Form, HTTPException, BackgroundTasks, Query
from fastapi.responses import StreamingResponse
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import get_settings
from app.core.database import get_db
from app.models.fleet_models import ImportJob, ImportStatus
from app.services.excel_import_service import import_excel
from app.api.fleet_common import FLEET_IMPORT
from app.utils.uploads import save_upload_limited

router = APIRouter(prefix="/fleet", tags=["Fleet - Import Excel"])
settings = get_settings()

@router.post("/phuong-tien/import", summary="Import Excel danh sách phương tiện",
             dependencies=[Depends(FLEET_IMPORT)])
async def import_phuong_tien(
    background_tasks: BackgroundTasks,
    file: UploadFile = File(...),
    header_row:    int  = Form(default=4, description="Dòng bắt đầu header (file Trường Phát = 4)"),
    data_start_row: int = Form(default=6, description="Dòng bắt đầu data (file Trường Phát = 6)"),
    preview: bool = Form(default=False, description="Preview: parse nhưng không lưu"),
    upsert_mode: str = Form(default="upsert", description="upsert|insert (insert=bỏ qua trùng)"),
    column_mapping: str = Form(default="", description="JSON {col_index: db_field} để override mapping"),
    db: AsyncSession = Depends(get_db),
):
    return await _handle_import(file, "phuong_tien", header_row, data_start_row, preview, background_tasks, db,
                                upsert_mode, column_mapping)

@router.post("/lai-xe/import", summary="Import Excel danh sách lái xe",
             dependencies=[Depends(FLEET_IMPORT)])
async def import_lai_xe(
    background_tasks: BackgroundTasks,
    file: UploadFile = File(...),
    header_row:    int  = Form(default=4),
    data_start_row: int = Form(default=6),
    preview: bool = Form(default=False, description="Preview: parse nhưng không lưu"),
    upsert_mode: str = Form(default="upsert"),
    column_mapping: str = Form(default="", description="JSON {col_index: db_field}"),
    db: AsyncSession = Depends(get_db),
):
    return await _handle_import(file, "lai_xe", header_row, data_start_row, preview, background_tasks, db,
                                upsert_mode, column_mapping)

async def _handle_import(file, table, header_row, data_start_row, preview, background_tasks, db,
                         upsert_mode: str = "upsert", column_mapping: str = ""):
    job_id = str(uuid.uuid4())
    filepath = settings.UPLOAD_DIR / f"{job_id}_{file.filename}"
    await save_upload_limited(file, filepath,
                              allowed_exts=(".xlsx", ".xls"),
                              max_mb=settings.MAX_UPLOAD_MB)

    mapping = {}
    if column_mapping:
        try:
            mapping = json.loads(column_mapping)
        except Exception:
            mapping = {}

    if preview:
        # Preview: parse trực tiếp, KHÔNG tạo ImportJob trong DB
        # (job tạo ở session này chưa commit nên session khác không thấy -> failed oan)
        result = await import_excel(db, str(filepath), table, f"preview-{job_id}",
                                    header_row, data_start_row, True,
                                    column_mapping=mapping, upsert_mode=upsert_mode)
        try:
            Path(filepath).unlink(missing_ok=True)
        except Exception:
            pass
        return {
            "mode": "preview",
            "total_rows":  result["total"],
            "valid_rows":  result["success"],
            "error_rows":  result["errors"],
            "sample_errors": result["error_details"][:10],
            "message": "Preview xong. Gọi lại với preview=false để import thật."
        }

    job = ImportJob(
        id=job_id,
        target_table=table,
        filename=file.filename,
        status=ImportStatus.PENDING,
        preview_mode=False,
        header_row=header_row,
        data_start_row=data_start_row,
        upsert_mode=upsert_mode if upsert_mode in ("upsert", "insert") else "upsert",
        column_mapping=column_mapping or None,
    )
    db.add(job)
    await db.flush()

    background_tasks.add_task(_run_import, job_id, str(filepath), table, header_row, data_start_row, False,
                              mapping, upsert_mode)
    return {
        "job_id": job_id,
        "status": "pending",
        "message": f"Import đang chạy background. Poll GET /fleet/import-jobs/{job_id}",
    }

@router.get("/import-jobs/{job_id}", summary="Xem tiến trình import")
async def get_import_job(job_id: str, db: AsyncSession = Depends(get_db)):
    result = await db.execute(select(ImportJob).where(ImportJob.id == job_id))
    job = result.scalar_one_or_none()
    if not job:
        raise HTTPException(404, "Import job không tồn tại")

    progress = 0
    if job.total_rows and job.total_rows > 0:
        progress = round(job.processed_rows / job.total_rows * 100, 1)

    return {
        "job_id": job.id,
        "table": job.target_table,
        "filename": job.filename,
        "status": job.status,
        "progress_percent": progress,
        "total_rows": job.total_rows,
        "success_rows": job.success_rows,
        "error_rows": job.error_rows,
        "created_at": job.created_at,
        "completed_at": job.completed_at,
        "error_details": json.loads(job.error_details) if job.error_details else [],
    }

@router.get("/import-jobs/{job_id}/errors", summary="Tải danh sách lỗi import")
async def get_import_errors(job_id: str, db: AsyncSession = Depends(get_db)):
    result = await db.execute(select(ImportJob).where(ImportJob.id == job_id))
    job = result.scalar_one_or_none()
    if not job:
        raise HTTPException(404, "Import job không tồn tại")

    errors = json.loads(job.error_details) if job.error_details else []
    return {"job_id": job_id, "total_errors": len(errors), "errors": errors}

@router.get("/import-jobs/{job_id}/errors-xlsx", summary="Tải file Excel các dòng lỗi (highlight)")
async def get_import_errors_xlsx(job_id: str, db: AsyncSession = Depends(get_db)):
    result = await db.execute(select(ImportJob).where(ImportJob.id == job_id))
    job = result.scalar_one_or_none()
    if not job:
        raise HTTPException(404, "Import job không tồn tại")
    errors = json.loads(job.error_details) if job.error_details else []
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Loi import"
    ws.append(["Dòng Excel", "Trường", "Lỗi"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="C0392B")
    for e in errors:
        ws.append([e.get("row"), e.get("field"), e.get("error")])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="loi_import_{job_id[:8]}.xlsx"'})

@router.post("/import/preview-headers", summary="Đọc header Excel để UI mapping cột",
             dependencies=[Depends(FLEET_IMPORT)])
async def preview_headers(
    file: UploadFile = File(...),
    header_row: int = Form(default=4),
    data_start_row: int = Form(default=6),
):
    import openpyxl
    from pathlib import Path as _Path
    import uuid as _uuid
    tmp_path = settings.TEMP_DIR / f"preview_{_uuid.uuid4().hex}.xlsx"
    await save_upload_limited(file, _Path(tmp_path),
                              allowed_exts=(".xlsx", ".xls"),
                              max_mb=settings.MAX_UPLOAD_MB)
    try:
        wb = openpyxl.load_workbook(tmp_path, read_only=True, data_only=True)
        ws = wb.active
        headers: list[str] = []
        for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if header_row <= row_idx < data_start_row:
                for v in row:
                    if v is not None and str(v).strip():
                        headers.append(str(v).strip())
        wb.close()
        # gợi ý field DB
        from app.services.excel_import_service import PHUONG_TIEN_MAP, LAI_XE_MAP
        suggest = {}
        for h in headers:
            k = h.lower().strip()
            if k in PHUONG_TIEN_MAP and PHUONG_TIEN_MAP[k]:
                suggest[h] = PHUONG_TIEN_MAP[k]
            elif k in LAI_XE_MAP and LAI_XE_MAP[k]:
                suggest[h] = LAI_XE_MAP[k]
        return {"headers": headers, "goi_y_mapping": suggest}
    finally:
        _Path(tmp_path).unlink(missing_ok=True)


# ─── Validate PATCH (chặn dữ liệu rác trước khi lưu) ──────────────────────
_PT_DATES = {"han_dang_kiem", "han_phu_hieu", "han_bao_hiem"}
_LX_DATES = {"han_gplx", "hop_dong_ngay_ky", "ksk_ngay_kham", "tap_huan_ngay"}

_PT_INTS = {"nam_san_xuat": (1900, 2100), "so_cho": (1, 100),
            "trong_tai_kg": (0, 100000), "so_ghe": (1, 100)}

async def _run_import(job_id: str, filepath: str, table: str, header_row: int, data_start_row: int, preview: bool = False,
                  column_mapping: dict | None = None, upsert_mode: str = "upsert"):
    from app.core.database import AsyncSessionLocal

    stats = {"total": 0, "success": 0, "errors": 0, "error_details": []}
    async with AsyncSessionLocal() as db:
        try:
            # Update status → processing
            result = await db.execute(select(ImportJob).where(ImportJob.id == job_id))
            job = result.scalar_one()
            job.status = ImportStatus.PROCESSING
            await db.commit()

            # Run import
            stats = await import_excel(db, filepath, table, job_id, header_row, data_start_row, preview,
                                       column_mapping=column_mapping, upsert_mode=upsert_mode)
            await db.commit()  # commit data import vào DB

            # Update job kết quả
            async with AsyncSessionLocal() as db2:
                result2 = await db2.execute(select(ImportJob).where(ImportJob.id == job_id))
                job2 = result2.scalar_one()
                job2.status = ImportStatus.DONE
                job2.total_rows = stats["total"]
                job2.success_rows = stats["success"]
                job2.error_rows = stats["errors"]
                job2.error_details = json.dumps(stats["error_details"], ensure_ascii=False) if stats["error_details"] else None
                job2.completed_at = datetime.now(timezone.utc)
                await db2.commit()

        except Exception as e:
            async with AsyncSessionLocal() as db_err:
                result3 = await db_err.execute(select(ImportJob).where(ImportJob.id == job_id))
                job3 = result3.scalar_one_or_none()
                if job3:
                    job3.status = ImportStatus.FAILED
                    job3.error_details = json.dumps([{"error": str(e)}])
                    job3.completed_at = datetime.now(timezone.utc)
                    await db_err.commit()

    return stats


# ─── Helpers ────────────────────────────────────────────────────────────────
